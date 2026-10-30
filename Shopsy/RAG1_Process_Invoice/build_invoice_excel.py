# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Parse the MarkItDown invoices in converted/ into data/invoices.xlsx.

Sheets:
  Invoices   - one row per invoice (header fields, totals, check status)
  Line Items - one row per invoice line, joined to invoice fields

MarkItDown's PDF tables are inconsistent (the same field may land in a
table cell or a plain line), so every line is flattened to plain text
before matching. Each invoice is then checked: qty x unit price = amount,
sum of lines = subtotal, subtotal + tax = total.

Run with:  uv run build_invoice_excel.py
"""
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
CONVERTED = BASE / "converted"
OUTPUT = BASE / "data" / "invoices.xlsx"

# Longest first so "AI Services + Expenses" wins over "AI Services".
SERVICE_TYPES = sorted(
    ["AI Services + Expenses", "AI Services", "AI Platform", "Cloud Expenses",
     "Software Licenses", "Travel & Accommodation"],
    key=len, reverse=True,
)
SEPARATOR = re.compile(r"^\|[\s|-]+\|$")
MONEY = r"\$([\d,]+\.\d{2})"
LINE_ITEM = re.compile(rf"^(.+?)\s+(\d+(?:\.\d+)?)\s+{MONEY}\s+{MONEY}$")


def money(s: str) -> float:
    return float(s.replace(",", ""))


def flatten(line: str) -> str:
    """Join markdown table cells into one space-separated string."""
    if line.startswith("|"):
        return " ".join(c.strip() for c in line.strip("|").split("|") if c.strip())
    return line.strip()


def cells(line: str) -> list[str]:
    if line.startswith("|"):
        return [c.strip() for c in line.strip("|").split("|") if c.strip()]
    return [line.strip()] if line.strip() else []


def parse_bill_to(raw_lines: list[str]) -> tuple[str, str, str]:
    """Return (client, address, service_type) from the BILL TO block."""
    parts = []
    for line in raw_lines:
        parts += [c for c in cells(line) if c not in ("BILL TO", "SERVICE TYPE")]

    service = ""
    rest = []
    for p in parts:
        if not service:
            if p in SERVICE_TYPES:
                service = p
                continue
            st = next((t for t in SERVICE_TYPES if p.endswith(" " + t)), None)
            if st:
                service = st
                p = p[: -len(st)].strip()
        rest.append(p)
    client = rest[0] if rest else ""
    address = ", ".join(rest[1:])
    return client, address, service


def parse_invoice(path: Path) -> tuple[dict, list[dict]]:
    raw = [l for l in path.read_text(encoding="utf-8").splitlines()
           if l.strip() and not SEPARATOR.match(l.strip())]
    flat = [flatten(l) for l in raw]
    text = "\n".join(flat)

    def find(pattern: str) -> str:
        m = re.search(pattern, text)
        if not m:
            raise ValueError(f"{path.name}: pattern not found: {pattern}")
        return m.group(1)

    inv_no = find(r"Invoice No\.\s+(\S+)")
    bill_start = next(i for i, l in enumerate(flat) if l.startswith("BILL TO"))
    desc_idx = next(i for i, l in enumerate(flat) if l.startswith("DESCRIPTION"))
    client, address, service = parse_bill_to(raw[bill_start:desc_idx])

    items = []
    for line in flat[desc_idx + 1:]:
        if line.startswith("Subtotal"):
            break
        m = LINE_ITEM.match(line)
        if not m:
            raise ValueError(f"{path.name}: unparsed line item: {line!r}")
        desc, qty, unit, amount = m.groups()
        items.append({"Description": desc, "Qty": float(qty),
                      "Unit Price": money(unit), "Amount": money(amount)})

    tax_rate, tax = re.search(rf"Tax \((\d+(?:\.\d+)?)%\):\s+{MONEY}", text).groups()
    notes_idx = flat.index("NOTES")
    invoice = {
        "Invoice No": inv_no,
        "Invoice Type": "Client Invoice" if inv_no.startswith("INV") else "Internal Expense",
        "Invoice Date": datetime.strptime(find(r"(?<!Due )Date:\s+(\d{4}-\d{2}-\d{2})"), "%Y-%m-%d"),
        "Due Date": datetime.strptime(find(r"Due Date:\s+(\d{4}-\d{2}-\d{2})"), "%Y-%m-%d"),
        "Client": client,
        "Client Address": address,
        "Service Type": service,
        "Line Count": len(items),
        "Subtotal": money(find(rf"Subtotal:\s+{MONEY}")),
        "Tax Rate": float(tax_rate) / 100,
        "Tax": money(tax),
        "Total Due": money(find(rf"TOTAL DUE\s+{MONEY}")),
        "Notes": flat[notes_idx + 1] if notes_idx + 1 < len(flat) else "",
        "Source File": path.name,
    }

    problems = []
    for it in items:
        if abs(it["Qty"] * it["Unit Price"] - it["Amount"]) > 0.01:
            problems.append(f"qty x price != amount ({it['Description']})")
    if abs(sum(it["Amount"] for it in items) - invoice["Subtotal"]) > 0.01:
        problems.append("lines != subtotal")
    if abs(invoice["Subtotal"] + invoice["Tax"] - invoice["Total Due"]) > 0.01:
        problems.append("subtotal + tax != total")
    if not (client and service):
        problems.append("missing client or service type")
    invoice["Check"] = "; ".join(problems) or "OK"

    return invoice, items


def write_sheet(ws, rows: list[dict], money_cols=(), pct_cols=(), date_cols=()):
    headers = list(rows[0])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[h] for h in headers])
    for idx, name in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        width = max(len(str(c.value or "")) for c in ws[letter])
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        fmt = ("#,##0.00" if name in money_cols else
               "0%" if name in pct_cols else
               "yyyy-mm-dd" if name in date_cols else None)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    files = sorted(CONVERTED.glob("*.md"))
    if not files:
        print(f"No markdown files in {CONVERTED}")
        return 1

    invoices, lines = [], []
    for f in files:
        inv, items = parse_invoice(f)
        d = inv["Invoice Date"]
        # Put Year/Month right after the date for pivoting.
        inv = dict(list(inv.items())[:3] + [("Year", d.year), ("Month", d.strftime("%Y-%m"))]
                   + list(inv.items())[3:])
        invoices.append(inv)
        for n, it in enumerate(items, start=1):
            lines.append({
                "Invoice No": inv["Invoice No"], "Line No": n,
                "Invoice Type": inv["Invoice Type"], "Invoice Date": inv["Invoice Date"],
                "Year": inv["Invoice Date"].year, "Month": inv["Invoice Date"].strftime("%Y-%m"),
                "Client": inv["Client"], "Service Type": inv["Service Type"], **it,
            })
        print(f"[{inv['Check']}] {inv['Invoice No']}: {inv['Client']} - ${inv['Total Due']:,.2f}")

    wb = Workbook()
    write_sheet(wb.active, invoices, money_cols=("Subtotal", "Tax", "Total Due"),
                pct_cols=("Tax Rate",), date_cols=("Invoice Date", "Due Date"))
    wb.active.title = "Invoices"
    write_sheet(wb.create_sheet("Line Items"), lines, money_cols=("Unit Price", "Amount"),
                date_cols=("Invoice Date",))
    OUTPUT.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT)

    bad = sum(inv["Check"] != "OK" for inv in invoices)
    print(f"\nWrote {OUTPUT} - {len(invoices)} invoices, {len(lines)} line items, "
          f"{bad} with check issues")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
