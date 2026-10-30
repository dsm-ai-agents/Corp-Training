import json, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
R = pathlib.Path(__file__).resolve().parent.parent
S = json.load(open(R/'working/kpi_summary.json')); T = S['totals']
NAVY, TEAL, ORANGE, GREY, LIGHT = RGBColor(0x14,0x21,0x3D), RGBColor(0x0F,0x76,0x6E), RGBColor(0xC2,0x41,0x0C), RGBColor(0x55,0x5F,0x6D), RGBColor(0xEE,0xF2,0xF6)
L = lambda v: f"Rs {v/1e5:,.2f} lakh"
p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[6])
def tb(x,y,w,h,paras,anchor=MSO_ANCHOR.TOP,name=None):
    b = s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=b.text_frame; f.word_wrap=True; f.vertical_anchor=anchor
    f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
    if name: b.name=name
    for i,(t,sz,bold,col) in enumerate(paras):
        para = f.paragraphs[0] if i==0 else f.add_paragraph()
        r = para.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold; r.font.color.rgb=col; r.font.name='Calibri'
        para.space_after = Pt(4)
    return b
top = S['zones'][0]
tb(0.5,0.35,12.3,0.9,[(f"Delivered GMV is {T['delivered_rate']*100:.0f}% of booked; {top['name'].title()} drives {top['share_of_delivered']*100:.0f}% of deliveries",30,True,NAVY)],name='Title')
tb(0.5,1.15,12.3,0.3,[(f"Regional sales performance, {S['meta']['period']}. Region = seller zone.",14,False,GREY)],name='Subtitle')
tiles = [('Booked GMV',T['booked_gmv'],f"{T['booked_orders']:,} booked orders",NAVY),('Delivered GMV',T['delivered_gmv'],f"{T['delivered_rate']*100:.1f}% of booked",TEAL),
 ('Returned GMV',T['returned_gmv'],f"{T['return_rate']*100:.1f}% of booked",ORANGE),('Cancelled GMV',T['cancelled_gmv'],f"{T['cancel_rate']*100:.1f}% of gross",ORANGE)]
for i,(lab,v,sub,col) in enumerate(tiles):
    x = 0.5+i*3.1
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.65),Inches(2.9),Inches(1.3)); r.fill.solid(); r.fill.fore_color.rgb=LIGHT; r.line.fill.background(); r.shadow.inherit=False; r.name=f'Tile {lab}'
    tb(x+0.2,1.75,2.5,1.1,[(lab,12,True,GREY),(L(v),24,True,col),(sub,12,False,GREY)],name=f'Tile text {lab}')
cd = CategoryChartData(); cd.categories=[f"{z['name'].title()} ({z['sellers']} sellers)" for z in S['zones']]
cd.add_series('Booked GMV (Rs lakh)',[round(z['booked_gmv']/1e5,2) for z in S['zones']])
cd.add_series('Delivered GMV (Rs lakh)',[round(z['delivered_gmv']/1e5,2) for z in S['zones']])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(0.5),Inches(3.25),Inches(7.4),Inches(3.4),cd); gf.name='Regional chart'
c = gf.chart; c.has_title=True; c.chart_title.text_frame.text='Booked vs delivered GMV by region (Rs lakh)'
tr = c.chart_title.text_frame.paragraphs[0].runs[0]; tr.font.size=Pt(14); tr.font.bold=True; tr.font.color.rgb=NAVY; tr.font.name='Calibri'
c.has_legend=True; c.legend.position=XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout=False; c.legend.font.size=Pt(12)
for ser,col in zip(c.plots[0].series,[NAVY,TEAL]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb=col
c.plots[0].has_data_labels=True; dl=c.plots[0].data_labels; dl.number_format='0.0'; dl.number_format_is_linked=False; dl.font.size=Pt(12); dl.position=XL_LABEL_POSITION.OUTSIDE_END
c.value_axis.visible=False; c.value_axis.has_major_gridlines=False; c.category_axis.tick_labels.font.size=Pt(12)
c.plots[0].gap_width=60
ex = S['exceptions']
paras = [('Sellers to act on',16,True,NAVY)]
for r in ex[:4]:
    fl = ', '.join(f.replace('Delivered rate 10+ pts below zone','low delivered rate').replace('High return rate','high returns').replace('High cancel rate','high cancels') for f in r['flags'])
    paras.append((f"{r['name'].replace('Training ','')} ({r['zone'].title()}, {r['manager']}): {fl}",14,False,GREY))
if len(ex)>4: paras.append((f"+{len(ex)-4} more flagged. Full list in the dashboard.",12,False,GREY))
tb(8.3,3.35,4.5,3.3,paras,name='Exceptions')
tb(0.5,6.85,12.3,0.4,[(f"Booked = Delivered + Ready to ship + Returned. Booked + Cancelled = Gross {L(T['gross_gmv'])}. {S['reconciliation']['checks_passed']} of {S['reconciliation']['checks_total']} data checks passed. Flag rule: return or cancel rate above 1.5x average, or delivered rate 10+ pts below zone.",10,False,GREY)],name='Footnote')
s.notes_slide.notes_text_frame.text = ' '.join(S['headlines'])
p.core_properties.title='Executive regional sales slide'
p.save(R/'outputs/executive_slide.pptx'); print('saved pptx')
